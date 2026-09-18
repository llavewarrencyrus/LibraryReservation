import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

books = [
    {
        "id": 1,
        "title": "The Great Gatsby",
        "author": "F. Scott Fitzgerald",
        "description": "A portrait of the Jazz Age, exploring themes of decadence, idealism, social upheaval, and the elusive American Dream through the mysterious Jay Gatsby.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780743273565-M.jpg",
        "section": "Circulation",
        "publisher": "Charles Scribner's Sons",
        "placeOfPublication": "New York, USA",
        "callNumber": "FIC FIT 1925",
        "status": "Available",
        "subjects": "Classic Literature, 1920s, American Fiction",
        "edition": "Scribner Paperback Edition"
    },
    {
        "id": 2,
        "title": "1984",
        "author": "George Orwell",
        "description": "A chilling dystopian novel set in Airstrip One, depicting the perils of totalitarianism, pervasive government surveillance, and thought control.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780451524935-M.jpg",
        "section": "Circulation",
        "publisher": "Secker & Warburg",
        "placeOfPublication": "London, UK",
        "callNumber": "FIC ORW 1949",
        "status": "Available",
        "subjects": "Dystopian, Political Fiction, Classics",
        "edition": "Signet Classic Edition"
    },
    {
        "id": 3,
        "title": "To Kill a Mockingbird",
        "author": "Harper Lee",
        "description": "The unforgettable story of childhood innocence, empathy, and racial injustice in Maycomb, Alabama, championed by attorney Atticus Finch.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780061120084-M.jpg",
        "section": "Circulation",
        "publisher": "J. B. Lippincott & Co.",
        "placeOfPublication": "Philadelphia, USA",
        "callNumber": "FIC LEE 1960",
        "status": "Available",
        "subjects": "Southern Gothic, Legal Drama, Bildungsroman",
        "edition": "Harper Perennial Modern Classics"
    },
    {
        "id": 4,
        "title": "Pride and Prejudice",
        "author": "Jane Austen",
        "description": "A timeless romantic comedy charting the emotional development of Elizabeth Bennet as she navigates manners, upbringing, morality, and marriage in Regency Britain.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780141439518-M.jpg",
        "section": "Circulation",
        "publisher": "T. Egerton",
        "placeOfPublication": "London, UK",
        "callNumber": "FIC AUS 1813",
        "status": "Available",
        "subjects": "Romance, British Literature, Regency Era",
        "edition": "Penguin Classics Deluxe"
    },
    {
        "id": 5,
        "title": "Moby-Dick",
        "author": "Herman Melville",
        "description": "The epic voyage of the whaling ship Pequod and Captain Ahab's monomaniacal vengeance against the giant white whale.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780142437247-M.jpg",
        "section": "Reference",
        "publisher": "Harper & Brothers",
        "placeOfPublication": "New York, USA",
        "callNumber": "FIC MEL 1851",
        "status": "Available",
        "subjects": "Adventure, Sea Stories, American Classics",
        "edition": "Penguin Classics Edition"
    },
    {
        "id": 6,
        "title": "War and Peace",
        "author": "Leo Tolstoy",
        "description": "A sweeping historical novel detailing the French invasion of Russia and the impact of the Napoleonic era on Tsarist society through five aristocratic families.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9781400079988-M.jpg",
        "section": "Reference",
        "publisher": "The Russian Messenger",
        "placeOfPublication": "Moscow, Russia",
        "callNumber": "FIC TOL 1869",
        "status": "Available",
        "subjects": "Historical Fiction, Russian Literature, War",
        "edition": "Vintage Classics Translation"
    },
    {
        "id": 7,
        "title": "The Catcher in the Rye",
        "author": "J.D. Salinger",
        "description": "A profound exploration of teenage alienation, identity, and the struggle to protect innocence, told from the iconic perspective of Holden Caulfield.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780316769488-M.jpg",
        "section": "Circulation",
        "publisher": "Little, Brown and Company",
        "placeOfPublication": "Boston, USA",
        "callNumber": "FIC SAL 1951",
        "status": "Available",
        "subjects": "Young Adult, Coming-of-Age, American Fiction",
        "edition": "Mass Market Paperback"
    },
    {
        "id": 8,
        "title": "The Odyssey",
        "author": "Homer",
        "description": "One of the foundation works of Western literature, recounting the ten-year journey of Odysseus returning home to Ithaca after the fall of Troy.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780140268866-M.jpg",
        "section": "Reference",
        "publisher": "Penguin Classics",
        "placeOfPublication": "London, UK",
        "callNumber": "883 HOM 1996",
        "status": "Available",
        "subjects": "Epic Poetry, Greek Mythology, Ancient History",
        "edition": "Robert Fagles Translation"
    },
    {
        "id": 9,
        "title": "The Brothers Karamazov",
        "author": "Fyodor Dostoevsky",
        "description": "A passionate philosophical novel delving into deep questions of faith, doubt, morality, free will, and patricide in 19th-century Russia.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780374528379-M.jpg",
        "section": "Circulation",
        "publisher": "The Russian Messenger",
        "placeOfPublication": "St. Petersburg, Russia",
        "callNumber": "FIC DOS 1880",
        "status": "Available",
        "subjects": "Philosophy, Psychological Fiction, Classics",
        "edition": "Farrar, Straus and Giroux Edition"
    },
    {
        "id": 10,
        "title": "Brave New World",
        "author": "Aldous Huxley",
        "description": "A visionary satire depicting a dystopian future where humanity is technologically engineered, chemically conditioned, and spiritually numb.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780060850524-M.jpg",
        "section": "Circulation",
        "publisher": "Chatto & Windus",
        "placeOfPublication": "London, UK",
        "callNumber": "FIC HUX 1932",
        "status": "Reserved",
        "subjects": "Science Fiction, Dystopian, Satire",
        "edition": "Harper Perennial Modern Classics"
    },
    {
        "id": 11,
        "title": "The Picture of Dorian Gray",
        "author": "Oscar Wilde",
        "description": "A gothic tale of aestheticism and moral decadence, wherein a handsome young man remains eternally youthful while his hidden portrait ages with every sin.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780141439570-M.jpg",
        "section": "Circulation",
        "publisher": "Ward, Lock and Company",
        "placeOfPublication": "London, UK",
        "callNumber": "FIC WIL 1890",
        "status": "Available",
        "subjects": "Gothic, Philosophical, Classic Literature",
        "edition": "Penguin Classics Edition"
    },
    {
        "id": 12,
        "title": "The Grapes of Wrath",
        "author": "John Steinbeck",
        "description": "The tragic journey of the Joad family, Oklahoma tenant farmers driven from their home during the Dust Bowl and migrating west to California.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780143039433-M.jpg",
        "section": "Circulation",
        "publisher": "The Viking Press",
        "placeOfPublication": "New York, USA",
        "callNumber": "FIC STE 1939",
        "status": "Available",
        "subjects": "Historical Fiction, Social Realism, Great Depression",
        "edition": "Viking Critical Edition"
    },
    {
        "id": 13,
        "title": "Jane Eyre",
        "author": "Charlotte Brontë",
        "description": "The passionate and independent journey of an orphaned governess who falls in love with the brooding Edward Rochester of Thornfield Hall.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780141441146-M.jpg",
        "section": "Circulation",
        "publisher": "Smith, Elder & Co.",
        "placeOfPublication": "London, UK",
        "callNumber": "FIC BRO 1847",
        "status": "Available",
        "subjects": "Gothic Romance, Victorian Literature, Feminism",
        "edition": "Penguin Classics Edition"
    },
    {
        "id": 14,
        "title": "Wuthering Heights",
        "author": "Emily Brontë",
        "description": "A tempestuous and brooding tale of intense, destructive passion and vengeance between Heathcliff and Catherine Earnshaw on the Yorkshire moors.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780141439556-M.jpg",
        "section": "Circulation",
        "publisher": "Thomas Cautley Newby",
        "placeOfPublication": "London, UK",
        "callNumber": "FIC BRO 1847b",
        "status": "Available",
        "subjects": "Gothic Fiction, Tragedy, British Classics",
        "edition": "Oxford World's Classics"
    },
    {
        "id": 15,
        "title": "Fahrenheit 451",
        "author": "Ray Bradbury",
        "description": "In a bleak dystopian society where books are outlawed and critical thought is suppressed, fireman Guy Montag begins to question his duty of burning literature.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9781451673319-M.jpg",
        "section": "Circulation",
        "publisher": "Ballantine Books",
        "placeOfPublication": "New York, USA",
        "callNumber": "FIC BRA 1953",
        "status": "Available",
        "subjects": "Dystopian, Science Fiction, Censorship",
        "edition": "60th Anniversary Edition"
    },
    {
        "id": 16,
        "title": "The Hobbit",
        "author": "J.R.R. Tolkien",
        "description": "Bilbo Baggins, a peace-loving hobbit, is swept into an epic quest by Gandalf and thirteen dwarves to reclaim the Lonely Mountain and its treasure from the dragon Smaug.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780547928227-M.jpg",
        "section": "Young Readers",
        "publisher": "George Allen & Unwin",
        "placeOfPublication": "London, UK",
        "callNumber": "FIC TOL 1937",
        "status": "Available",
        "subjects": "High Fantasy, Adventure, Middle-earth",
        "edition": "75th Anniversary Illustrated Edition"
    },
    {
        "id": 17,
        "title": "The Fellowship of the Ring",
        "author": "J.R.R. Tolkien",
        "description": "Frodo Baggins inherits the One Ring of Sauron and sets out on a perilous journey across Middle-earth accompanied by an alliance of heroes.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780544003415-M.jpg",
        "section": "Circulation",
        "publisher": "George Allen & Unwin",
        "placeOfPublication": "London, UK",
        "callNumber": "FIC TOL 1954",
        "status": "Available",
        "subjects": "High Fantasy, Epic, Adventure",
        "edition": "Deluxe Hardcover Edition"
    },
    {
        "id": 18,
        "title": "The Two Towers",
        "author": "J.R.R. Tolkien",
        "description": "The fellowship fractured, Frodo and Sam journey toward Mordor with Gollum while Aragorn, Legolas, and Gimli face battles across Rohan.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780547928203-M.jpg",
        "section": "Circulation",
        "publisher": "George Allen & Unwin",
        "placeOfPublication": "London, UK",
        "callNumber": "FIC TOL 1954b",
        "status": "Available",
        "subjects": "High Fantasy, Epic, Middle-earth",
        "edition": "Mariner Books Edition"
    },
    {
        "id": 19,
        "title": "The Return of the King",
        "author": "J.R.R. Tolkien",
        "description": "The climatic finale of the War of the Ring as the armies of Gondor rally for the final confrontation and Frodo approaches Mount Doom.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780547928197-M.jpg",
        "section": "Circulation",
        "publisher": "George Allen & Unwin",
        "placeOfPublication": "London, UK",
        "callNumber": "FIC TOL 1955",
        "status": "Unavailable",
        "subjects": "High Fantasy, Epic, Climax",
        "edition": "Mariner Books Edition"
    },
    {
        "id": 20,
        "title": "The Lion, the Witch and the Wardrobe",
        "author": "C.S. Lewis",
        "description": "Four siblings step through a magic wardrobe into the snowy kingdom of Narnia, trapped in eternal winter by the tyrannical White Witch.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780064404990-M.jpg",
        "section": "Children's Section",
        "publisher": "Geoffrey Bles",
        "placeOfPublication": "London, UK",
        "callNumber": "JUV LEW 1950",
        "status": "Available",
        "subjects": "Children's Fantasy, Allegory, Adventure",
        "edition": "HarperCollins Children's Edition"
    },
    {
        "id": 21,
        "title": "The Alchemist",
        "author": "Paulo Coelho",
        "description": "An enchanting mystical fable about Santiago, an Andalusian shepherd boy who travels to the Egyptian pyramids in search of treasure and his personal legend.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780062315007-M.jpg",
        "section": "Circulation",
        "publisher": "HarperOne",
        "placeOfPublication": "New York, USA",
        "callNumber": "FIC COE 1988",
        "status": "Available",
        "subjects": "Philosophical Fiction, Quest, Inspiration",
        "edition": "25th Anniversary Edition"
    },
    {
        "id": 22,
        "title": "The Fault in Our Stars",
        "author": "John Green",
        "description": "A heart-wrenching and witty romance between two teenage cancer patients, Hazel Grace Lancaster and Augustus Waters, navigating love and mortality.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780525478812-M.jpg",
        "section": "Young Readers",
        "publisher": "Dutton Books",
        "placeOfPublication": "New York, USA",
        "callNumber": "YA GRE 2012",
        "status": "Available",
        "subjects": "Young Adult, Contemporary Romance, Drama",
        "edition": "First Edition Hardcover"
    },
    {
        "id": 23,
        "title": "The Kite Runner",
        "author": "Khaled Hosseini",
        "description": "A powerful and emotional story of friendship, betrayal, and redemption set against the turbulent background of Afghanistan over three decades.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9781594631931-M.jpg",
        "section": "Circulation",
        "publisher": "Riverhead Books",
        "placeOfPublication": "New York, USA",
        "callNumber": "FIC HOS 2003",
        "status": "Available",
        "subjects": "Historical Drama, Friendship, Middle Eastern Literature",
        "edition": "10th Anniversary Edition"
    },
    {
        "id": 24,
        "title": "The Book Thief",
        "author": "Markus Zusak",
        "description": "Narrated by Death, this poignant novel tells the story of Liesel Meminger, a foster girl living outside of Munich who scratches out an existence by stealing books.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780375842207-M.jpg",
        "section": "Circulation",
        "publisher": "Knopf Books for Young Readers",
        "placeOfPublication": "New York, USA",
        "callNumber": "FIC ZUS 2005",
        "status": "Available",
        "subjects": "Historical Fiction, WWII, Holocaust",
        "edition": "Alfred A. Knopf Hardcover"
    },
    {
        "id": 25,
        "title": "Life of Pi",
        "author": "Yann Martel",
        "description": "A remarkable tale of survival, faith, and zoology following Piscine Molitor 'Pi' Patel, stranded on a lifeboat in the Pacific Ocean with a Bengal tiger.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780156027328-M.jpg",
        "section": "Circulation",
        "publisher": "Knopf Canada",
        "placeOfPublication": "Toronto, Canada",
        "callNumber": "FIC MAR 2001",
        "status": "Available",
        "subjects": "Adventure Fiction, Philosophical, Survival",
        "edition": "Harvest Books Paperback"
    },
    {
        "id": 26,
        "title": "The Road",
        "author": "Cormac McCarthy",
        "description": "A devastating and lyrical post-apocalyptic journey of an unnamed father and son walking through a burned, ash-covered American landscape toward the coast.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780307387899-M.jpg",
        "section": "Circulation",
        "publisher": "Alfred A. Knopf",
        "placeOfPublication": "New York, USA",
        "callNumber": "FIC MCC 2006",
        "status": "Available",
        "subjects": "Post-Apocalyptic, Fatherhood, Modern Classics",
        "edition": "Vintage Books Paperback"
    },
    {
        "id": 27,
        "title": "The Handmaid's Tale",
        "author": "Margaret Atwood",
        "description": "Set in the Republic of Gilead, a religious totalitarian regime where fertile women are subjugated as child-bearing handmaids.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780385490818-M.jpg",
        "section": "Circulation",
        "publisher": "McClelland and Stewart",
        "placeOfPublication": "Toronto, Canada",
        "callNumber": "FIC ATW 1985",
        "status": "Available",
        "subjects": "Dystopian, Feminist Fiction, Speculative",
        "edition": "Anchor Books Edition"
    },
    {
        "id": 28,
        "title": "The Bell Jar",
        "author": "Sylvia Plath",
        "description": "A haunting semi-autobiographical chronicle of Esther Greenwood's descent into clinical depression amidst the societal pressures of 1950s America.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780060837020-M.jpg",
        "section": "Circulation",
        "publisher": "Heinemann",
        "placeOfPublication": "London, UK",
        "callNumber": "FIC PLA 1963",
        "status": "Available",
        "subjects": "Psychological Fiction, Mental Health, Classics",
        "edition": "Harper Perennial Edition"
    },
    {
        "id": 29,
        "title": "The Secret Garden",
        "author": "Frances Hodgson Burnett",
        "description": "Orphaned Mary Lennox discovers an overgrown walled garden at Misselthwaite Manor and brings healing, friendship, and joy to her cousin Colin.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780141321066-M.jpg",
        "section": "Children's Section",
        "publisher": "Frederick A. Stokes",
        "placeOfPublication": "New York, USA",
        "callNumber": "JUV BUR 1911",
        "status": "Available",
        "subjects": "Children's Literature, Nature, Friendship",
        "edition": "Puffin Classics Edition"
    },
    {
        "id": 30,
        "title": "The Little Prince",
        "author": "Antoine de Saint-Exupéry",
        "description": "A poetic and philosophical tale of an aviator stranded in the Sahara desert who meets an ethereal prince fallen to Earth from a tiny asteroid.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780156012195-M.jpg",
        "section": "Children's Section",
        "publisher": "Reynal & Hitchcock",
        "placeOfPublication": "New York, USA",
        "callNumber": "FIC SAI 1943",
        "status": "Available",
        "subjects": "Fable, Philosophy, French Literature",
        "edition": "Harcourt Brace Translation"
    },
    {
        "id": 31,
        "title": "Great Expectations",
        "author": "Charles Dickens",
        "description": "The dramatic transformation of orphan Pip, aided by an anonymous benefactor, exploring wealth, class, remorse, and moral redemption.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780141439563-M.jpg",
        "section": "Circulation",
        "publisher": "Chapman & Hall",
        "placeOfPublication": "London, UK",
        "callNumber": "FIC DIC 1861",
        "status": "Available",
        "subjects": "Victorian Literature, Bildungsroman, Social Satire",
        "edition": "Penguin Classics Revised"
    },
    {
        "id": 32,
        "title": "The Old Man and the Sea",
        "author": "Ernest Hemingway",
        "description": "The heroic struggle between aging Cuban fisherman Santiago and a giant marlin far out in the Gulf Stream waters.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780684801223-M.jpg",
        "section": "Circulation",
        "publisher": "Charles Scribner's Sons",
        "placeOfPublication": "New York, USA",
        "callNumber": "FIC HEM 1952",
        "status": "Available",
        "subjects": "Novella, Courage, American Literature",
        "edition": "Scribner Paperback"
    },
    {
        "id": 33,
        "title": "The Metamorphosis",
        "author": "Franz Kafka",
        "description": "Traveling salesman Gregor Samsa awakens one morning to find himself inexplicably transformed into a monstrous vermin.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780553213690-M.jpg",
        "section": "Circulation",
        "publisher": "Kurt Wolff Verlag",
        "placeOfPublication": "Leipzig, Germany",
        "callNumber": "FIC KAF 1915",
        "status": "Available",
        "subjects": "Absurdist Fiction, Existentialism, Novella",
        "edition": "Bantam Classic Translation"
    },
    {
        "id": 34,
        "title": "The Hitchhiker's Guide to the Galaxy",
        "author": "Douglas Adams",
        "description": "Minutes before Earth is demolished for a hyperspace bypass, Arthur Dent is whisked into space on an absurd interstellar romp.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780345391803-M.jpg",
        "section": "Circulation",
        "publisher": "Pan Books",
        "placeOfPublication": "London, UK",
        "callNumber": "FIC ADA 1979",
        "status": "Available",
        "subjects": "Comic Science Fiction, Space, Humor",
        "edition": "Del Rey Mass Market Paperback"
    },
    {
        "id": 35,
        "title": "The Color Purple",
        "author": "Alice Walker",
        "description": "An epistolary novel tracing the life, suffering, and ultimate spiritual triumph of Celie, a poor African American woman in early 20th-century rural Georgia.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780156028356-M.jpg",
        "section": "Circulation",
        "publisher": "Harcourt Brace Jovanovich",
        "placeOfPublication": "New York, USA",
        "callNumber": "FIC WAL 1982",
        "status": "Available",
        "subjects": "Epistolary Novel, African American, Feminist",
        "edition": "Harvest Books Edition"
    },
    {
        "id": 36,
        "title": "The Art of War",
        "author": "Sun Tzu",
        "description": "Ancient Chinese military treatise offering timeless strategies on leadership, tactical maneuver, psychology, and victory without direct conflict.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9781590302255-M.jpg",
        "section": "Reference",
        "publisher": "Shambhala Publications",
        "placeOfPublication": "Boston, USA",
        "callNumber": "355.02 SUN 2005",
        "status": "Available",
        "subjects": "Military Strategy, Philosophy, Eastern Thought",
        "edition": "Thomas Cleary Translation"
    },
    {
        "id": 37,
        "title": "The Diary of a Young Girl",
        "author": "Anne Frank",
        "description": "The candid and inspiring diary written by a Jewish girl hiding with her family in a secret annex in Amsterdam during the Nazi occupation.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780553296983-M.jpg",
        "section": "Reference",
        "publisher": "Contact Publishing",
        "placeOfPublication": "Amsterdam, Netherlands",
        "callNumber": "940.53 FRA 1947",
        "status": "Available",
        "subjects": "Autobiography, WWII, Holocaust History",
        "edition": "Definitive Edition"
    },
    {
        "id": 38,
        "title": "Catch-22",
        "author": "Joseph Heller",
        "description": "A satirical masterpiece following US Army Air Forces bombardier Yossarian and his struggle to preserve his sanity amidst bureaucratic wartime madness.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9781451626650-M.jpg",
        "section": "Circulation",
        "publisher": "Simon & Schuster",
        "placeOfPublication": "New York, USA",
        "callNumber": "FIC HEL 1961",
        "status": "Available",
        "subjects": "Satire, War Fiction, Black Comedy",
        "edition": "50th Anniversary Edition"
    },
    {
        "id": 39,
        "title": "The Giver",
        "author": "Lois Lowry",
        "description": "Twelve-year-old Jonas lives in a seemingly utopian world devoid of suffering or color until he is chosen to be the Receiver of Memory.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780544336261-M.jpg",
        "section": "Young Readers",
        "publisher": "Houghton Mifflin",
        "placeOfPublication": "Boston, USA",
        "callNumber": "YA LOW 1993",
        "status": "Available",
        "subjects": "Dystopian, YA Literature, Newbery Medal",
        "edition": "20th Anniversary Edition"
    },
    {
        "id": 40,
        "title": "Noli Me Tangere",
        "author": "José Rizal",
        "description": "The passionate national novel exposing the abuses and hypocrisy of the Spanish friars and colonial administration in the Philippines through Crisostomo Ibarra.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780143039693-M.jpg",
        "section": "Filipiniana",
        "publisher": "Berliner Buchdruckerei-Aktien-Gesellschaft",
        "placeOfPublication": "Berlin, Germany",
        "callNumber": "FIL FIC RIZ 1887",
        "status": "Available",
        "subjects": "Philippine Revolution, Nationalism, Social Reform",
        "edition": "Harold Augenbraum Translation"
    },
    {
        "id": 41,
        "title": "El Filibusterismo",
        "author": "José Rizal",
        "description": "The darker sequel to Noli Me Tangere, following the vengeful jeweler Simoun as he schemes an armed insurrection against Spanish rule.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780143106395-M.jpg",
        "section": "Filipiniana",
        "publisher": "F. Meyer-van Loo Press",
        "placeOfPublication": "Ghent, Belgium",
        "callNumber": "FIL FIC RIZ 1891",
        "status": "Available",
        "subjects": "Filipiniana, Political Fiction, Anti-Colonialism",
        "edition": "Penguin Classics Edition"
    },
    {
        "id": 42,
        "title": "Florante at Laura",
        "author": "Francisco Balagtas",
        "description": "A masterwork of Philippine Tagalog literature and metrical romance narrating love, courage, injustice, and religious tolerance in the kingdom of Albania.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9789710848034-M.jpg",
        "section": "Filipiniana",
        "publisher": "Colegio de Santo Tomas",
        "placeOfPublication": "Manila, Philippines",
        "callNumber": "FIL 899.211 BAL 1838",
        "status": "Available",
        "subjects": "Philippine Poetry, Epic, Awit",
        "edition": "National Bookstore Edition"
    },
    {
        "id": 43,
        "title": "Sapiens: A Brief History of Humankind",
        "author": "Yuval Noah Harari",
        "description": "A provocative narrative examining how biology and history have defined Homo sapiens, from the Cognitive Revolution to contemporary biotechnology.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780062316097-M.jpg",
        "section": "Non-Fiction",
        "publisher": "Harper",
        "placeOfPublication": "New York, USA",
        "callNumber": "909 HAR 2014",
        "status": "Available",
        "subjects": "Anthropology, World History, Evolutionary Biology",
        "edition": "HarperCollins Paperback"
    },
    {
        "id": 44,
        "title": "Thinking, Fast and Slow",
        "author": "Daniel Kahneman",
        "description": "Nobel laureate Daniel Kahneman explains the two systems of human cognition: fast, intuitive thinking and slow, deliberate deliberation.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780374533557-M.jpg",
        "section": "Non-Fiction",
        "publisher": "Farrar, Straus and Giroux",
        "placeOfPublication": "New York, USA",
        "callNumber": "153.4 KAH 2011",
        "status": "Reserved",
        "subjects": "Cognitive Psychology, Behavioral Economics, Decision Making",
        "edition": "First Paperbound Edition"
    },
    {
        "id": 45,
        "title": "Atomic Habits",
        "author": "James Clear",
        "description": "A practical framework for breaking bad habits, building good routines, and achieving remarkable results through continuous 1% daily improvements.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780735211292-M.jpg",
        "section": "Non-Fiction",
        "publisher": "Avery / Penguin Random House",
        "placeOfPublication": "New York, USA",
        "callNumber": "158.1 CLE 2018",
        "status": "Available",
        "subjects": "Self-Improvement, Habit Formation, Productivity",
        "edition": "Trade Hardcover"
    },
    {
        "id": 46,
        "title": "The Midnight Library",
        "author": "Matt Haig",
        "description": "Between life and death there is a mystical library with infinite shelves, giving Nora Seed the chance to undo past regrets and choose another life.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780525559474-M.jpg",
        "section": "Circulation",
        "publisher": "Viking",
        "placeOfPublication": "London, UK",
        "callNumber": "FIC HAI 2020",
        "status": "Available",
        "subjects": "Fantasy, Magical Realism, Contemporary Fiction",
        "edition": "First Edition Hardcover"
    },
    {
        "id": 47,
        "title": "Where the Crawdads Sing",
        "author": "Delia Owens",
        "description": "An exquisite ode to the natural world and a murder mystery following Kya Clark, the 'Marsh Girl' who raised herself in the wilderness of North Carolina.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780735219090-M.jpg",
        "section": "Circulation",
        "publisher": "G.P. Putnam's Sons",
        "placeOfPublication": "New York, USA",
        "callNumber": "FIC OWE 2018",
        "status": "Available",
        "subjects": "Coming-of-Age, Mystery, Nature Writing",
        "edition": "First Edition Paperback"
    },
    {
        "id": 48,
        "title": "Circe",
        "author": "Madeline Miller",
        "description": "A vivid and feminist reimagining of the life of Circe, the daughter of Helios, banished to the island of Aiaia where she hones her occult powers.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780316556347-M.jpg",
        "section": "Circulation",
        "publisher": "Little, Brown and Company",
        "placeOfPublication": "Boston, USA",
        "callNumber": "FIC MIL 2018",
        "status": "Available",
        "subjects": "Mythology, Fantasy, Retelling",
        "edition": "Little Brown Hardcover"
    },
    {
        "id": 49,
        "title": "Dune",
        "author": "Frank Herbert",
        "description": "Set on the desert planet Arrakis, young Paul Atreides must navigate political intrigue, ecology, religion, and spice to claim his destiny.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780441013593-M.jpg",
        "section": "Circulation",
        "publisher": "Chilton Books",
        "placeOfPublication": "Philadelphia, USA",
        "callNumber": "FIC HER 1965",
        "status": "Available",
        "subjects": "Science Fiction, Space Opera, Ecology",
        "edition": "Ace Books Premium Edition"
    },
    {
        "id": 50,
        "title": "A Brief History of Time",
        "author": "Stephen Hawking",
        "description": "A landmark non-fiction book on theoretical cosmology, explaining black holes, the Big Bang, general relativity, and the search for a unified theory.",
        "imageUrl": "https://covers.openlibrary.org/b/isbn/9780553380163-M.jpg",
        "section": "Science & Technology",
        "publisher": "Bantam Books",
        "placeOfPublication": "New York, USA",
        "callNumber": "523.1 HAW 1988",
        "status": "Available",
        "subjects": "Astrophysics, Cosmology, Popular Science",
        "edition": "Updated and Expanded Edition"
    }
]

headers = [
    "id",
    "title",
    "author",
    "description",
    "imageUrl",
    "section",
    "publisher",
    "placeOfPublication",
    "callNumber",
    "status",
    "subjects",
    "edition"
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "BookList"

# Styling definitions
header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Navy Blue
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

thin_border = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
normal_font = Font(name="Calibri", size=10)

# Write headers
ws.append(headers)

for col_num in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = header_alignment
    cell.border = thin_border

ws.row_dimensions[1].height = 26

# Write rows
for row_idx, book in enumerate(books, start=2):
    row_data = [book.get(h, "") for h in headers]
    ws.append(row_data)
    
    use_zebra = (row_idx % 2 == 0)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.font = normal_font
        cell.border = thin_border
        if use_zebra:
            cell.fill = zebra_fill
            
        header_name = headers[col_idx - 1]
        if header_name in ["id"]:
            cell.alignment = Alignment(horizontal="center", vertical="top")
        elif header_name in ["status", "callNumber", "section"]:
            cell.alignment = Alignment(horizontal="center", vertical="top")
        elif header_name in ["description", "subjects"]:
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        else:
            cell.alignment = Alignment(horizontal="left", vertical="top")
            
    ws.row_dimensions[row_idx].height = 36

# Set column widths
col_widths = {
    "A": 8,   # id
    "B": 32,  # title
    "C": 24,  # author
    "D": 50,  # description
    "E": 30,  # imageUrl
    "F": 18,  # section
    "G": 28,  # publisher
    "H": 22,  # placeOfPublication
    "I": 18,  # callNumber
    "J": 14,  # status
    "K": 32,  # subjects
    "L": 26   # edition
}

for col_letter, width in col_widths.items():
    ws.column_dimensions[col_letter].width = width

# Freeze top header row
ws.freeze_panes = "A2"

output_path = "booklist.xlsx"
wb.save(output_path)
print(f"Successfully generated {output_path} with {len(books)} books.")
